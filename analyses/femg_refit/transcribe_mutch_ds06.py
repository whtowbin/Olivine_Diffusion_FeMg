"""Write the olivine regressions of Mutch et al. (2021) Data Set S6 to a CSV.

Mutch, Maclennan, Shorttle, Rudge and Neave (2021), DFENS: Diffusion chronometry using
Finite Elements and Nested Sampling. G-cubed 22(4), doi:10.1029/2020GC009303; Data Set
S6 (`data/mutch2021_dfens_ds06.xlsx`, as distributed with the DFENS repository). The
same regressions as Mutch et al. (2019) Supplementary Tables 1-2, at full precision:

    ln D[001] = a + b ln fO2(bar) + c X_Fo + q/T + j P + h P/T     (P in Pa, T in K)

"Version 1" (with h; the one DFENS uses), fO2 in bar (the authors' recommended form).
One row per law and parameter: the coefficient, then its row of the covariance.

Run: uv run --with openpyxl python analyses/femg_refit/transcribe_mutch_ds06.py
Writes data/mutch2021_dfens_ds06_olivine.csv.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).parent
XLSX = HERE / "data" / "mutch2021_dfens_ds06.xlsx"
OUT = HERE / "data" / "mutch2021_dfens_ds06_olivine.csv"
PARAMS = ("a", "b", "c", "q", "j", "h")
# Sheet rows (1-based) of each law: coefficients on "Olivine Regression Parameters",
# the first covariance row on "Olivine Regression Covariance " (Version 1, bar).
LAWS = {"FeMg (Global)": (8, 9), "FeMg (TaMED)": (9, 16), "Ni": (10, 23), "Mn": (11, 30)}


def main() -> None:
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    par, cov = wb["Olivine Regression Parameters"], wb["Olivine Regression Covariance "]
    rows = []
    for law, (prow, crow) in LAWS.items():
        assert str(par.cell(prow, 1).value).strip() == law, par.cell(prow, 1).value
        assert str(cov.cell(crow - 1, 1).value).strip() == law, cov.cell(crow - 1, 1).value
        for i, p in enumerate(PARAMS):
            assert par.cell(7, 2 + i).value == p and cov.cell(crow + i, 1).value == p
            row = {"law": law, "param": p, "value": float(par.cell(prow, 2 + i).value)}
            row |= {
                f"cov_{q}": float(cov.cell(crow + i, 2 + k).value) for k, q in enumerate(PARAMS)
            }
            rows.append(row)
    pd.DataFrame(rows).to_csv(OUT, index=False, float_format="%.15g")
    print(f"{len(rows)} rows -> {OUT}")


if __name__ == "__main__":
    main()
